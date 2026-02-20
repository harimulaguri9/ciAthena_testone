from openpyxl import load_workbook
import json

class ExcelReader:
    def __init__(self, file_path, sheet_name):
        self.wb = load_workbook(file_path)
        self.sheet = self.wb[sheet_name]

    def get_test_data(self):
        data = []
        for row in self.sheet.iter_rows(min_row=2, values_only=True):
            question = row[0]
            expected_sql = row[1]
            expected_be_response = row[2]   # JSON stored in Excel
            chart = row[3]

            if question:
                data.append({
                    "question": question,
                    "expected_sql": expected_sql,
                    "expected_be_response": expected_be_response,
                    "chart": chart
                })
        return data