from openpyxl import Workbook

class ExcelWriter:
    def __init__(self, file_path):
        self.file_path = file_path
        self.wb = Workbook()
        self.sheet = self.wb.active
        self.row = 1

        headers = [
            "Question",
            "SQL_Match",
            "BE_Response_Match",
            "Chart_Match",
            "Stream_Completed",
            "STATUS"
        ]

        for col, header in enumerate(headers, start=1):
            self.sheet.cell(row=1, column=col, value=header)

    def write_row(self, question, sql_match, be_match, chart_match, stream_status, status):
        self.row += 1
        self.sheet.cell(row=self.row, column=1, value=question)
        self.sheet.cell(row=self.row, column=2, value=sql_match)
        self.sheet.cell(row=self.row, column=3, value=be_match)
        self.sheet.cell(row=self.row, column=4, value=chart_match)
        self.sheet.cell(row=self.row, column=5, value=stream_status)
        self.sheet.cell(row=self.row, column=6, value=status)

    def save(self):
        self.wb.save(self.file_path)